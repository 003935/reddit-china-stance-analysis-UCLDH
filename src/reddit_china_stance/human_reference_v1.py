"""Fail-closed import of the v1 human semantic-reference workbooks.

The source workbooks are never modified.  Row text, context, labels, review
notes, and row-level hashes remain below the ignored ``data/`` tree.  Public
outputs are aggregate, metadata-only receipts.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import shutil
import unicodedata
from collections import Counter, defaultdict
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

from jsonschema import Draft202012Validator
from openpyxl import load_workbook

from reddit_china_stance.human_seeded_consensus_v1 import canonical_sha256
from reddit_china_stance.privacy import assert_metadata_only

REPO_ROOT = Path(__file__).resolve().parents[2]
SCHEMA_VERSION = "1.0.0"
KIND = "human-reference-semantic-v1"

DEFAULT_SEMANTIC_WORKBOOK = Path("data/semantic-labels.xlsx")
DEFAULT_REVIEW_WORKBOOK = Path("data/human-review.xlsx")
DEFAULT_SOURCE_PACKET = REPO_ROOT / "data/private-gemma-review/source-packet.parquet"
DEFAULT_SCHEMA_PATH = REPO_ROOT / "schemas/human-reference-semantic-label-v1.schema.json"
PRIVATE_ROOT = REPO_ROOT / "data/private-human-reference-v1"
PUBLIC_ROOT = REPO_ROOT / "outputs/human-reference-v1"

SEMANTIC_SHEET = "Annotations"
SEMANTIC_HEADER_ROW = 3
SEMANTIC_HEADERS = (
    "Sample ID",
    "TARGET TEXT — label this",
    "Submission context — reference only",
    "Parent context — reference only",
    "Truncation notice",
    "Language",
    "Relevance",
    "Target 1",
    "Stance 1",
    "Target 2",
    "Stance 2",
    "Target 3",
    "Stance 3",
    "Target 4",
    "Stance 4",
    "Needs review?",
    "Review reason",
    "Annotator note",
    "QC status",
    "QC detail",
)
REVIEW_SHEET = "Review"
REVIEW_HEADER_ROW = 1
REVIEW_HEADERS = (
    "Review ID",
    "Target text — label this",
    "Submission context",
    "Parent context",
    "Your label",
    "Context used",
    "Short reason (optional)",
    "Truncation note",
)

LANGUAGES = frozenset(
    {"confident_english", "mixed_english_chinese", "chinese", "other", "unclear"}
)
RELEVANCE = frozenset({"material", "not_material", "unclear"})
REVIEW_RELEVANCE_MAP = {
    "relevant": "material",
    "not relevant": "not_material",
    "unclear": "unclear",
}
CONTEXT_USED = frozenset({"none", "submission", "parent", "both"})

NEAR_OVERLAP_THRESHOLD = 0.90
NEAR_OVERLAP_MIN_CHARS = 30
OVERLAP_FIELDS = ("target_text", "submission_context", "parent_context", "surface")
OVERLAP_PUBLIC_NAMES = {
    "target_text": "target",
    "submission_context": "submission",
    "parent_context": "parent",
    "surface": "combined_surface",
}
QUALITATIVE_AUDIT_FIRST_ROW = 4
QUALITATIVE_AUDIT_LAST_ROW = 103
DEFAULT_EXPECTED_SPLIT_COUNTS = {
    "development": 222,
    "locked_test_candidate": 230,
    "excluded_language": 36,
}

PRIVATE_ARTIFACT_NAMES = {
    "reference_rows": "reference-rows.json",
    "quarantine_rows": "quarantine-rows.json",
    "review_rows": "review-rows.json",
    "overlap_details": "overlap-details.json",
}


class HumanReferenceError(RuntimeError):
    """Raised when an input or immutable artefact violates the frozen contract."""


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while block := handle.read(1024 * 1024):
            digest.update(block)
    return digest.hexdigest()


def _file_descriptor(path: Path) -> dict[str, Any]:
    if not path.is_file():
        raise HumanReferenceError(f"required input is not a file: {path}")
    return {"name": path.name, "sha256": _sha256_file(path), "bytes": path.stat().st_size}


def _repo_bound_descriptor(path: Path) -> dict[str, Any]:
    descriptor = _file_descriptor(path)
    try:
        relative_path = path.resolve().relative_to(REPO_ROOT.resolve())
    except ValueError:
        relative_path = Path(path.name)
    return {**descriptor, "relative_path": str(relative_path)}


def _canonical_json_bytes(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n").encode()


def _write_immutable_json(path: Path, value: Mapping[str, Any]) -> dict[str, Any]:
    payload = _canonical_json_bytes(value)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        if path.read_bytes() != payload:
            raise HumanReferenceError(f"immutable JSON output differs: {path}")
    else:
        temporary = path.with_suffix(path.suffix + ".incomplete")
        if temporary.exists():
            raise HumanReferenceError(f"stale incomplete output exists: {temporary}")
        temporary.write_bytes(payload)
        os.replace(temporary, path)
    return {
        "relative_path": str(path.relative_to(REPO_ROOT)),
        "sha256": _sha256_file(path),
        "bytes": path.stat().st_size,
    }


def _copy_immutable(source: Path, destination: Path) -> dict[str, Any]:
    expected = _sha256_file(source)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if _sha256_file(destination) != expected:
            raise HumanReferenceError(f"immutable source copy differs: {destination}")
    else:
        temporary = destination.with_suffix(destination.suffix + ".incomplete")
        if temporary.exists():
            raise HumanReferenceError(f"stale incomplete source copy exists: {temporary}")
        shutil.copyfile(source, temporary)
        if _sha256_file(temporary) != expected:
            raise HumanReferenceError("source copy digest changed during copying")
        os.replace(temporary, destination)
    return {
        "relative_path": str(destination.relative_to(REPO_ROOT)),
        "sha256": expected,
        "bytes": destination.stat().st_size,
    }


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise HumanReferenceError(f"expected JSON object: {path}")
    return value


def _optional_text(value: Any, *, field: str) -> str | None:
    if value is None:
        return None
    if not isinstance(value, str):
        raise HumanReferenceError(f"{field} must be text or blank")
    return value if value != "" else None


def _text_or_none_for_quarantine(value: Any) -> str | None:
    if value is None:
        return None
    return value if isinstance(value, str) else str(value)


def _header_values(sheet: Any, *, row: int, columns: int) -> tuple[Any, ...]:
    return tuple(sheet.cell(row=row, column=column).value for column in range(1, columns + 1))


def _open_exact_sheet(path: Path, *, sheet_name: str, header_row: int, headers: tuple[str, ...]):
    workbook = load_workbook(path, read_only=True, data_only=False)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise HumanReferenceError(f"{path.name} is missing exact sheet {sheet_name!r}")
    sheet = workbook[sheet_name]
    actual = _header_values(sheet, row=header_row, columns=len(headers))
    if actual != headers:
        workbook.close()
        raise HumanReferenceError(f"{path.name} {sheet_name!r} headers differ from the contract")
    return workbook, sheet


def load_label_schema(path: Path = DEFAULT_SCHEMA_PATH) -> dict[str, Any]:
    schema = _read_json(path)
    Draft202012Validator.check_schema(schema)
    expected_top = {"relevance", "target_stances"}
    if set(schema.get("required", [])) != expected_top:
        raise HumanReferenceError("human-reference schema required fields drifted")
    return schema


def _target_orders(schema: Mapping[str, Any]) -> tuple[dict[str, int], set[str], set[str]]:
    item_properties = schema["properties"]["target_stances"]["items"]["properties"]
    target_values = list(item_properties["target"]["enum"])
    stance_values = set(item_properties["stance"]["enum"])
    return (
        {value: index for index, value in enumerate(target_values)},
        set(target_values),
        stance_values,
    )


def _semantic_reason_codes(
    raw: Mapping[str, Any],
    *,
    duplicate_sample_id: bool,
    schema: Mapping[str, Any],
) -> tuple[list[str], dict[str, Any] | None]:
    reasons: set[str] = set()
    sample_id = raw[SEMANTIC_HEADERS[0]]
    target_text = raw[SEMANTIC_HEADERS[1]]
    language = raw[SEMANTIC_HEADERS[5]]
    relevance = raw[SEMANTIC_HEADERS[6]]
    if not isinstance(sample_id, str) or not sample_id:
        reasons.add("missing_or_invalid_sample_id")
    if duplicate_sample_id:
        reasons.add("duplicate_sample_id")
    if not isinstance(target_text, str) or not target_text:
        reasons.add("missing_or_invalid_target_text")
    if language is None or language == "":
        reasons.add("missing_language")
    elif language not in LANGUAGES:
        reasons.add("invalid_language")
    if relevance is None or relevance == "":
        reasons.add("missing_relevance")
    elif relevance not in RELEVANCE:
        reasons.add("invalid_relevance")

    target_order, targets, stances = _target_orders(schema)
    target_stances: list[dict[str, str]] = []
    seen_empty = False
    for slot in range(4):
        target = raw[SEMANTIC_HEADERS[7 + slot * 2]]
        stance = raw[SEMANTIC_HEADERS[8 + slot * 2]]
        target_missing = target is None or target == ""
        stance_missing = stance is None or stance == ""
        if target_missing and stance_missing:
            seen_empty = True
            continue
        if seen_empty:
            reasons.add("noncontiguous_target_stance_slots")
        if target_missing != stance_missing:
            reasons.add(f"partial_target_stance_pair_{slot + 1}")
        if not target_missing and target not in targets:
            reasons.add(f"invalid_target_{slot + 1}")
        if not stance_missing and stance not in stances:
            reasons.add(f"invalid_stance_{slot + 1}")
        if not target_missing and not stance_missing and target in targets and stance in stances:
            target_stances.append({"target": target, "stance": stance})
    target_values = [item["target"] for item in target_stances]
    if len(set(target_values)) != len(target_values):
        reasons.add("duplicate_target")
    any_target_value = any(
        raw[SEMANTIC_HEADERS[7 + slot * 2]] not in (None, "") for slot in range(4)
    )
    any_stance_value = any(
        raw[SEMANTIC_HEADERS[8 + slot * 2]] not in (None, "") for slot in range(4)
    )
    if relevance == "material" and not target_stances:
        reasons.add("material_without_complete_target_stance")
    if relevance in {"not_material", "unclear"} and (any_target_value or any_stance_value):
        reasons.add("nonmaterial_with_target_stance")

    if reasons:
        return sorted(reasons), None
    label = {
        "relevance": relevance,
        "target_stances": sorted(target_stances, key=lambda item: target_order[item["target"]]),
    }
    errors = sorted(
        Draft202012Validator(schema).iter_errors(label), key=lambda error: list(error.path)
    )
    if errors:
        return ["schema_validation_failed"], None
    if relevance == "material" and not label["target_stances"]:
        return ["material_without_complete_target_stance"], None
    if relevance != "material" and label["target_stances"]:
        return ["nonmaterial_with_target_stance"], None
    return [], label


def _raw_mapping(headers: tuple[str, ...], values: Sequence[Any]) -> dict[str, Any]:
    return {
        header: values[index] if index < len(values) else None
        for index, header in enumerate(headers)
    }


def _load_source_packet(path: Path, *, expected_rows: int) -> tuple[set[str], dict[str, Any]]:
    try:
        import pyarrow.parquet as pq
    except ImportError as exc:  # pragma: no cover - repository runtime normally supplies pyarrow
        raise HumanReferenceError("pyarrow is required to read the frozen source packet") from exc
    table = pq.read_table(path, columns=["record_id", "target_text"])
    if table.num_rows != expected_rows:
        raise HumanReferenceError(
            f"source packet has {table.num_rows} rows; expected exactly {expected_rows}"
        )
    rows = table.to_pylist()
    if any(not isinstance(row["record_id"], str) or not row["record_id"] for row in rows):
        raise HumanReferenceError("source packet contains an invalid record_id")
    if len({row["record_id"] for row in rows}) != len(rows):
        raise HumanReferenceError("source packet record_ids are not unique")
    if any(not isinstance(row["target_text"], str) or not row["target_text"] for row in rows):
        raise HumanReferenceError("source packet contains an invalid target_text")
    targets = {row["target_text"] for row in rows}
    return targets, {
        **_file_descriptor(path),
        "rows": table.num_rows,
        "unique_target_texts": len(targets),
    }


def _semantic_rows(
    path: Path,
    *,
    schema: Mapping[str, Any],
    source_packet_targets: set[str],
    expected_rows: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook, sheet = _open_exact_sheet(
        path,
        sheet_name=SEMANTIC_SHEET,
        header_row=SEMANTIC_HEADER_ROW,
        headers=SEMANTIC_HEADERS,
    )
    try:
        source_rows: list[tuple[int, dict[str, Any]]] = []
        for row_index, cells in enumerate(
            sheet.iter_rows(
                min_row=SEMANTIC_HEADER_ROW + 1,
                min_col=1,
                max_col=len(SEMANTIC_HEADERS),
                values_only=True,
            ),
            start=SEMANTIC_HEADER_ROW + 1,
        ):
            raw = _raw_mapping(SEMANTIC_HEADERS, cells)
            if any(raw[header] not in (None, "") for header in SEMANTIC_HEADERS[:18]):
                source_rows.append((row_index, raw))
    finally:
        workbook.close()
    if len(source_rows) != expected_rows:
        raise HumanReferenceError(
            f"semantic workbook has {len(source_rows)} data rows; expected exactly {expected_rows}"
        )
    counts = Counter(raw[SEMANTIC_HEADERS[0]] for _, raw in source_rows)
    valid_rows: list[dict[str, Any]] = []
    quarantine: list[dict[str, Any]] = []
    workbook_sha256 = _sha256_file(path)
    for source_row, raw in source_rows:
        sample_id = raw[SEMANTIC_HEADERS[0]]
        reasons, label = _semantic_reason_codes(
            raw,
            duplicate_sample_id=sample_id not in (None, "") and counts[sample_id] > 1,
            schema=schema,
        )
        raw_private = {
            key: _text_or_none_for_quarantine(raw[key]) for key in SEMANTIC_HEADERS[:18]
        }
        if reasons:
            body = {
                "source_kind": "semantic_workbook",
                "source_row": source_row,
                "source_sample_id": _text_or_none_for_quarantine(sample_id),
                "reason_codes": reasons,
                "raw_values": raw_private,
            }
            quarantine.append({**body, "row_sha256": canonical_sha256(body)})
            continue
        assert label is not None
        target_text = _optional_text(raw[SEMANTIC_HEADERS[1]], field="target_text")
        assert target_text is not None
        submission_context = _optional_text(
            raw[SEMANTIC_HEADERS[2]], field="submission_context"
        )
        parent_context = _optional_text(raw[SEMANTIC_HEADERS[3]], field="parent_context")
        truncation_notice = _optional_text(
            raw[SEMANTIC_HEADERS[4]], field="truncation_notice"
        )
        language = str(raw[SEMANTIC_HEADERS[5]])
        exposure_reasons = (
            ["prior_machine_packet"] if target_text in source_packet_targets else []
        )
        split = (
            "development"
            if language == "confident_english" and exposure_reasons
            else "locked_test_candidate"
            if language == "confident_english"
            else "excluded_language"
        )
        surface = {
            "target_text": target_text,
            "submission_context": submission_context,
            "parent_context": parent_context,
        }
        reference = {"language": language, "label": label}
        body = {
            "source_workbook_sha256": workbook_sha256,
            "source_sample_id": str(sample_id),
            "source_row": source_row,
            "split": split,
            "exposure_reasons": exposure_reasons,
            "language": language,
            **surface,
            "truncation_notice": truncation_notice,
            "label": label,
            "surface_sha256": canonical_sha256(surface),
            "reference_sha256": canonical_sha256(reference),
        }
        valid_rows.append({**body, "row_sha256": canonical_sha256(body)})
    if len({row["row_sha256"] for row in valid_rows}) != len(valid_rows):
        raise HumanReferenceError("semantic reference row hashes are not unique")
    return valid_rows, quarantine


def _review_rows(
    path: Path,
    *,
    expected_rows: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], set[str]]:
    workbook, sheet = _open_exact_sheet(
        path,
        sheet_name=REVIEW_SHEET,
        header_row=REVIEW_HEADER_ROW,
        headers=REVIEW_HEADERS,
    )
    try:
        source_rows: list[tuple[int, dict[str, Any]]] = []
        for row_index, cells in enumerate(
            sheet.iter_rows(
                min_row=REVIEW_HEADER_ROW + 1,
                min_col=1,
                max_col=len(REVIEW_HEADERS),
                values_only=True,
            ),
            start=REVIEW_HEADER_ROW + 1,
        ):
            raw = _raw_mapping(REVIEW_HEADERS, cells)
            if any(value not in (None, "") for value in raw.values()):
                source_rows.append((row_index, raw))
    finally:
        workbook.close()
    if len(source_rows) != expected_rows:
        raise HumanReferenceError(
            f"review workbook has {len(source_rows)} data rows; expected exactly {expected_rows}"
        )
    review_id_counts = Counter(raw[REVIEW_HEADERS[0]] for _, raw in source_rows)
    valid: list[dict[str, Any]] = []
    quarantine: list[dict[str, Any]] = []
    exposure_surfaces: set[str] = set()
    workbook_sha256 = _sha256_file(path)
    for source_row, raw in source_rows:
        reasons: set[str] = set()
        review_id = raw[REVIEW_HEADERS[0]]
        target_text = raw[REVIEW_HEADERS[1]]
        review_label = raw[REVIEW_HEADERS[4]]
        context_used = raw[REVIEW_HEADERS[5]]
        if not isinstance(review_id, str) or not review_id:
            reasons.add("missing_or_invalid_review_id")
        elif review_id_counts[review_id] > 1:
            reasons.add("duplicate_review_id")
        if not isinstance(target_text, str) or not target_text:
            reasons.add("missing_or_invalid_target_text")
        if review_label not in REVIEW_RELEVANCE_MAP:
            reasons.add("missing_or_invalid_review_relevance")
        if context_used not in CONTEXT_USED:
            reasons.add("missing_or_invalid_context_used")
        for header in REVIEW_HEADERS[2:4] + REVIEW_HEADERS[6:8]:
            value = raw[header]
            if value is not None and not isinstance(value, str):
                reasons.add(f"invalid_text_field_{REVIEW_HEADERS.index(header) + 1}")
        raw_private = {
            key: _text_or_none_for_quarantine(raw[key]) for key in REVIEW_HEADERS
        }
        valid_surface = (
            isinstance(target_text, str)
            and bool(target_text)
            and all(
                raw[header] is None or isinstance(raw[header], str)
                for header in REVIEW_HEADERS[2:4]
            )
        )
        if valid_surface:
            exposure_surfaces.add(
                canonical_sha256(
                    {
                        "target_text": target_text,
                        "submission_context": raw[REVIEW_HEADERS[2]] or None,
                        "parent_context": raw[REVIEW_HEADERS[3]] or None,
                    }
                )
            )
        if reasons:
            body = {
                "source_kind": "review_workbook",
                "source_row": source_row,
                "source_review_id": _text_or_none_for_quarantine(review_id),
                "reason_codes": sorted(reasons),
                "raw_values": raw_private,
            }
            quarantine.append({**body, "row_sha256": canonical_sha256(body)})
            continue
        surface = {
            "target_text": target_text,
            "submission_context": _optional_text(
                raw[REVIEW_HEADERS[2]], field="review_submission_context"
            ),
            "parent_context": _optional_text(
                raw[REVIEW_HEADERS[3]], field="review_parent_context"
            ),
        }
        surface_sha256 = canonical_sha256(surface)
        body = {
            "source_workbook_sha256": workbook_sha256,
            "source_review_id": str(review_id),
            "source_row": source_row,
            **surface,
            "relevance": REVIEW_RELEVANCE_MAP[str(review_label)],
            "context_used": str(context_used),
            "short_reason": _optional_text(raw[REVIEW_HEADERS[6]], field="short_reason"),
            "truncation_note": _optional_text(
                raw[REVIEW_HEADERS[7]], field="review_truncation_note"
            ),
            "surface_sha256": surface_sha256,
            "semantic_row_sha256s": [],
        }
        valid.append({**body, "row_sha256": canonical_sha256(body)})
    if len({row["row_sha256"] for row in valid}) != len(valid):
        raise HumanReferenceError("review row hashes are not unique")
    return valid, quarantine, exposure_surfaces


def _apply_exposure_policy(
    semantic_rows: Sequence[Mapping[str, Any]],
    *,
    review_surface_sha256s: set[str],
) -> list[dict[str, Any]]:
    updated: list[dict[str, Any]] = []
    for source in semantic_rows:
        row = dict(source)
        row.pop("row_sha256", None)
        reasons = set(row.get("exposure_reasons", []))
        if row["surface_sha256"] in review_surface_sha256s:
            reasons.add("review50_overlap")
        if QUALITATIVE_AUDIT_FIRST_ROW <= int(row["source_row"]) <= QUALITATIVE_AUDIT_LAST_ROW:
            reasons.add("qualitative_first100_audit")
        row["exposure_reasons"] = sorted(reasons)
        if row["language"] == "confident_english":
            row["split"] = "development" if reasons else "locked_test_candidate"
        else:
            row["split"] = "excluded_language"
        updated.append({**row, "row_sha256": canonical_sha256(row)})
    if len({row["row_sha256"] for row in updated}) != len(updated):
        raise HumanReferenceError("post-exposure semantic row hashes are not unique")
    return updated


def _link_review_rows(
    review_rows: Sequence[Mapping[str, Any]],
    *,
    semantic_rows: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    semantic_by_surface: defaultdict[str, list[str]] = defaultdict(list)
    for semantic in semantic_rows:
        semantic_by_surface[str(semantic["surface_sha256"])].append(
            str(semantic["row_sha256"])
        )
    linked: list[dict[str, Any]] = []
    for source in review_rows:
        row = dict(source)
        row.pop("row_sha256", None)
        row["semantic_row_sha256s"] = sorted(semantic_by_surface[str(row["surface_sha256"])])
        linked.append({**row, "row_sha256": canonical_sha256(row)})
    if len({row["row_sha256"] for row in linked}) != len(linked):
        raise HumanReferenceError("linked review row hashes are not unique")
    return linked


def _normalise_overlap_text(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).casefold().split())


def _surface_overlap_value(row: Mapping[str, Any], field: str) -> str | None:
    if field == "surface":
        parts = [row.get(name) for name in OVERLAP_FIELDS[:3]]
        if not any(parts):
            return None
        return "\n<context-boundary>\n".join(
            part if isinstance(part, str) else "" for part in parts
        )
    value = row.get(field)
    return value if isinstance(value, str) and value else None


def _word_shingles(value: str) -> frozenset[str]:
    normalised = _normalise_overlap_text(value)
    words = re.findall(r"\w+", normalised, flags=re.UNICODE)
    if len(words) >= 3:
        return frozenset(" ".join(words[index : index + 3]) for index in range(len(words) - 2))
    if len(normalised) >= 5:
        return frozenset(normalised[index : index + 5] for index in range(len(normalised) - 4))
    return frozenset({normalised}) if normalised else frozenset()


def _jaccard(left: frozenset[str], right: frozenset[str]) -> float:
    if not left or not right:
        return 0.0
    return len(left & right) / len(left | right)


def _cross_split_overlap(
    development: Sequence[Mapping[str, Any]],
    locked: Sequence[Mapping[str, Any]],
    *,
    field: str,
) -> dict[str, Any]:
    exact_pairs: list[dict[str, Any]] = []
    near_pairs: list[dict[str, Any]] = []
    max_nonexact = 0.0
    left_values = [(_surface_overlap_value(row, field), row) for row in development]
    right_values = [(_surface_overlap_value(row, field), row) for row in locked]
    right_exact: defaultdict[str, list[Mapping[str, Any]]] = defaultdict(list)
    for value, row in right_values:
        if value:
            right_exact[value].append(row)
    for value, row in left_values:
        if value:
            for other in right_exact[value]:
                exact_pairs.append(
                    {
                        "development_row_sha256": row["row_sha256"],
                        "locked_row_sha256": other["row_sha256"],
                    }
                )
    prepared_left = [
        (value, _word_shingles(value), row)
        for value, row in left_values
        if value and len(_normalise_overlap_text(value)) >= NEAR_OVERLAP_MIN_CHARS
    ]
    prepared_right = [
        (value, _word_shingles(value), row)
        for value, row in right_values
        if value and len(_normalise_overlap_text(value)) >= NEAR_OVERLAP_MIN_CHARS
    ]
    for left_value, left_shingles, left_row in prepared_left:
        for right_value, right_shingles, right_row in prepared_right:
            if left_value == right_value:
                continue
            similarity = _jaccard(left_shingles, right_shingles)
            max_nonexact = max(max_nonexact, similarity)
            if similarity >= NEAR_OVERLAP_THRESHOLD:
                near_pairs.append(
                    {
                        "development_row_sha256": left_row["row_sha256"],
                        "locked_row_sha256": right_row["row_sha256"],
                        "similarity": round(similarity, 6),
                    }
                )
    return {
        "exact_pairs": exact_pairs,
        "near_pairs": near_pairs,
        "max_nonexact_similarity": round(max_nonexact, 6),
    }


def _within_split_exact_groups(
    rows: Sequence[Mapping[str, Any]], *, field: str
) -> list[dict[str, Any]]:
    groups: defaultdict[str, list[str]] = defaultdict(list)
    for row in rows:
        value = _surface_overlap_value(row, field)
        if value:
            groups[hashlib.sha256(value.encode()).hexdigest()].append(str(row["row_sha256"]))
    return [
        {"value_sha256": value_sha256, "row_sha256s": sorted(row_hashes)}
        for value_sha256, row_hashes in sorted(groups.items())
        if len(row_hashes) > 1
    ]


def _overlap_details(rows: Sequence[Mapping[str, Any]]) -> dict[str, Any]:
    development = [row for row in rows if row["split"] == "development"]
    locked = [row for row in rows if row["split"] == "locked_test_candidate"]
    fields: dict[str, Any] = {}
    for field in OVERLAP_FIELDS:
        cross = _cross_split_overlap(development, locked, field=field)
        within = {
            "development": _within_split_exact_groups(development, field=field),
            "locked_test_candidate": _within_split_exact_groups(locked, field=field),
        }
        fields[field] = {"cross_split": cross, "within_split_exact_groups": within}
    return {
        "schema_version": SCHEMA_VERSION,
        "kind": f"{KIND}-private-overlap-details",
        "method": {
            "exact": "byte-identical non-empty field value",
            "near": "NFKC/casefold/whitespace-normalised word-trigram Jaccard",
            "near_threshold": NEAR_OVERLAP_THRESHOLD,
            "near_minimum_normalised_characters": NEAR_OVERLAP_MIN_CHARS,
        },
        "fields": fields,
    }


def _overlap_summary(details: Mapping[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field, value in details["fields"].items():
        cross = value["cross_split"]
        within = value["within_split_exact_groups"]
        result[OVERLAP_PUBLIC_NAMES[field]] = {
            "cross_split_exact_pairs": len(cross["exact_pairs"]),
            "cross_split_near_pairs": len(cross["near_pairs"]),
            "cross_split_max_nonexact_similarity": cross["max_nonexact_similarity"],
            "development_exact_duplicate_groups": len(within["development"]),
            "development_max_exact_group_size": max(
                (len(group["row_sha256s"]) for group in within["development"]), default=1
            ),
            "locked_test_exact_duplicate_groups": len(within["locked_test_candidate"]),
            "locked_test_max_exact_group_size": max(
                (
                    len(group["row_sha256s"])
                    for group in within["locked_test_candidate"]
                ),
                default=1,
            ),
        }
    return result


def make_contract(
    *,
    semantic_workbook: Path = DEFAULT_SEMANTIC_WORKBOOK,
    review_workbook: Path = DEFAULT_REVIEW_WORKBOOK,
    source_packet: Path = DEFAULT_SOURCE_PACKET,
    schema_path: Path = DEFAULT_SCHEMA_PATH,
    expected_semantic_rows: int = 500,
    expected_review_rows: int = 50,
    expected_source_packet_rows: int = 150,
    expected_split_counts: Mapping[str, int] | None = DEFAULT_EXPECTED_SPLIT_COUNTS,
) -> dict[str, Any]:
    if min(expected_semantic_rows, expected_review_rows, expected_source_packet_rows) <= 0:
        raise ValueError("expected row counts must be positive")
    expected_splits = (
        None
        if expected_split_counts is None
        else {
            key: int(expected_split_counts[key])
            for key in ("development", "locked_test_candidate", "excluded_language")
        }
    )
    if expected_splits is not None and (
        any(value < 0 for value in expected_splits.values())
        or sum(expected_splits.values()) > expected_semantic_rows
    ):
        raise ValueError("expected split counts are invalid")
    sources = {
        "semantic_workbook": _file_descriptor(semantic_workbook),
        "review_workbook": _file_descriptor(review_workbook),
        "source_packet": {
            **_file_descriptor(source_packet),
            "expected_rows": expected_source_packet_rows,
        },
        "label_schema": _repo_bound_descriptor(schema_path),
        "importer": _file_descriptor(Path(__file__).resolve()),
    }
    return {
        "schema_version": SCHEMA_VERSION,
        "kind": KIND,
        "sources": sources,
        "workbook_contract": {
            "semantic": {
                "sheet": SEMANTIC_SHEET,
                "header_row": SEMANTIC_HEADER_ROW,
                "headers": list(SEMANTIC_HEADERS),
                "expected_rows": expected_semantic_rows,
            },
            "review": {
                "sheet": REVIEW_SHEET,
                "header_row": REVIEW_HEADER_ROW,
                "headers": list(REVIEW_HEADERS),
                "expected_rows": expected_review_rows,
            },
        },
        "split_policy": {
            "development": (
                "structurally valid confident-English row with a recorded prior exposure"
            ),
            "locked_test_candidate": (
                "structurally valid confident-English row with no recorded prior exposure"
            ),
            "excluded_language": "structurally valid but not confident-English",
            "quarantine": "any structural or schema failure",
            "test_labels_inspected_for_selection": False,
            "exposure_reasons": [
                "prior_machine_packet",
                "review50_overlap",
                "qualitative_first100_audit",
            ],
            "qualitative_audit_source_rows": [
                QUALITATIVE_AUDIT_FIRST_ROW,
                QUALITATIVE_AUDIT_LAST_ROW,
            ],
            "expected_split_counts": expected_splits,
        },
        "overlap_policy": {
            "fields": list(OVERLAP_FIELDS),
            "near_threshold": NEAR_OVERLAP_THRESHOLD,
            "near_minimum_normalised_characters": NEAR_OVERLAP_MIN_CHARS,
        },
        "scientific_status": {
            "semantic_reference": "single_human_capability_reference",
            "review_reference": "single_human_relevance_bridge_only",
            "population_representative": False,
            "final_thesis_validation": False,
        },
        "privacy": {
            "private_rows_under_ignored_data": True,
            "public_receipts_metadata_only": True,
            "authors_read": False,
            "source_workbooks_immutable": True,
        },
    }


def manifest_path_for_run(run_id: str) -> Path:
    return PRIVATE_ROOT / f"manifest-{run_id}.json"


def private_run_root(run_id: str) -> Path:
    return PRIVATE_ROOT / f"run={run_id}"


def public_run_root(run_id: str) -> Path:
    return PUBLIC_ROOT / f"run={run_id}"


def manifest_path(
    *,
    semantic_workbook: Path = DEFAULT_SEMANTIC_WORKBOOK,
    review_workbook: Path = DEFAULT_REVIEW_WORKBOOK,
    source_packet: Path = DEFAULT_SOURCE_PACKET,
    schema_path: Path = DEFAULT_SCHEMA_PATH,
    expected_semantic_rows: int = 500,
    expected_review_rows: int = 50,
    expected_source_packet_rows: int = 150,
    expected_split_counts: Mapping[str, int] | None = DEFAULT_EXPECTED_SPLIT_COUNTS,
) -> Path:
    contract = make_contract(
        semantic_workbook=semantic_workbook,
        review_workbook=review_workbook,
        source_packet=source_packet,
        schema_path=schema_path,
        expected_semantic_rows=expected_semantic_rows,
        expected_review_rows=expected_review_rows,
        expected_source_packet_rows=expected_source_packet_rows,
        expected_split_counts=expected_split_counts,
    )
    return manifest_path_for_run(canonical_sha256(contract))


def _artifact_value(
    *, kind: str, run_id: str, rows: Sequence[Mapping[str, Any]]
) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "kind": kind,
        "run_id_value": run_id,
        "rows": list(rows),
    }


def prepare(
    *,
    semantic_workbook: Path = DEFAULT_SEMANTIC_WORKBOOK,
    review_workbook: Path = DEFAULT_REVIEW_WORKBOOK,
    source_packet: Path = DEFAULT_SOURCE_PACKET,
    schema_path: Path = DEFAULT_SCHEMA_PATH,
    expected_semantic_rows: int = 500,
    expected_review_rows: int = 50,
    expected_source_packet_rows: int = 150,
    expected_split_counts: Mapping[str, int] | None = DEFAULT_EXPECTED_SPLIT_COUNTS,
) -> dict[str, Any]:
    contract = make_contract(
        semantic_workbook=semantic_workbook,
        review_workbook=review_workbook,
        source_packet=source_packet,
        schema_path=schema_path,
        expected_semantic_rows=expected_semantic_rows,
        expected_review_rows=expected_review_rows,
        expected_source_packet_rows=expected_source_packet_rows,
        expected_split_counts=expected_split_counts,
    )
    run_id = canonical_sha256(contract)
    run_root = private_run_root(run_id)
    schema = load_label_schema(schema_path)
    packet_targets, packet_descriptor = _load_source_packet(
        source_packet, expected_rows=expected_source_packet_rows
    )
    if packet_descriptor["sha256"] != contract["sources"]["source_packet"]["sha256"]:
        raise HumanReferenceError("source packet changed after contract creation")
    semantic_rows, semantic_quarantine = _semantic_rows(
        semantic_workbook,
        schema=schema,
        source_packet_targets=packet_targets,
        expected_rows=expected_semantic_rows,
    )
    review_rows, review_quarantine, review_exposure_surfaces = _review_rows(
        review_workbook,
        expected_rows=expected_review_rows,
    )
    semantic_rows = _apply_exposure_policy(
        semantic_rows, review_surface_sha256s=review_exposure_surfaces
    )
    review_rows = _link_review_rows(review_rows, semantic_rows=semantic_rows)
    overlap_details = _overlap_details(semantic_rows)
    quarantine_rows = [*semantic_quarantine, *review_quarantine]

    split_counts = Counter(row["split"] for row in semantic_rows)
    actual_split_counts = {
        key: split_counts[key]
        for key in ("development", "locked_test_candidate", "excluded_language")
    }
    if expected_split_counts is not None and actual_split_counts != {
        key: int(expected_split_counts[key]) for key in actual_split_counts
    }:
        raise HumanReferenceError(
            f"exposure-aware split counts differ: {actual_split_counts}"
        )
    reason_counts = Counter(
        reason for row in quarantine_rows for reason in row.get("reason_codes", [])
    )
    summary = {
        "semantic_source_rows": expected_semantic_rows,
        "semantic_valid_rows": len(semantic_rows),
        "review_source_rows": expected_review_rows,
        "review_valid_rows": len(review_rows),
        "quarantine_rows": len(quarantine_rows),
        "split_counts": actual_split_counts,
        "quarantine_reason_counts": dict(sorted(reason_counts.items())),
        "review_exact_surface_matches": sum(
            bool(row["semantic_row_sha256s"]) for row in review_rows
        ),
        "review_exposure_surface_matches": sum(
            row["surface_sha256"] in review_exposure_surfaces for row in semantic_rows
        ),
    }

    private_artifacts: dict[str, dict[str, Any]] = {}
    artefact_rows = {
        "reference_rows": _artifact_value(
            kind=f"{KIND}-private-reference-rows", run_id=run_id, rows=semantic_rows
        ),
        "quarantine_rows": _artifact_value(
            kind=f"{KIND}-private-quarantine-rows", run_id=run_id, rows=quarantine_rows
        ),
        "review_rows": _artifact_value(
            kind=f"{KIND}-private-review-rows", run_id=run_id, rows=review_rows
        ),
    }
    for key, value in artefact_rows.items():
        descriptor = _write_immutable_json(run_root / PRIVATE_ARTIFACT_NAMES[key], value)
        private_artifacts[key] = {**descriptor, "rows": len(value["rows"])}
    overlap_descriptor = _write_immutable_json(
        run_root / PRIVATE_ARTIFACT_NAMES["overlap_details"], overlap_details
    )
    private_artifacts["overlap_details"] = overlap_descriptor

    source_copies = {
        "semantic_workbook": _copy_immutable(
            semantic_workbook,
            run_root
            / "source"
            / f"semantic-{contract['sources']['semantic_workbook']['sha256']}.xlsx",
        ),
        "review_workbook": _copy_immutable(
            review_workbook,
            run_root / "source" / f"review-{contract['sources']['review_workbook']['sha256']}.xlsx",
        ),
        "source_packet": _copy_immutable(
            source_packet,
            run_root
            / "source"
            / f"source-packet-{contract['sources']['source_packet']['sha256']}.parquet",
        ),
    }
    body = {
        "contract": contract,
        "run_id_value": run_id,
        "private_artifacts": private_artifacts,
        "source_copies": source_copies,
        "summary": summary,
        "overlap_summary": _overlap_summary(overlap_details),
    }
    manifest = {**body, "manifest_id": canonical_sha256(body)}
    manifest_descriptor = _write_immutable_json(manifest_path_for_run(run_id), manifest)
    receipt_body = {
        "schema_version": SCHEMA_VERSION,
        "kind": f"{KIND}-prepare-receipt",
        "status": "prepared",
        "run_id_value": run_id,
        "manifest_id": manifest["manifest_id"],
        "manifest_sha256": manifest_descriptor["sha256"],
        "source_hashes": {
            key: value["sha256"] for key, value in contract["sources"].items()
        },
        "counts": summary,
        "overlap_summary": manifest["overlap_summary"],
        "scientific_status": contract["scientific_status"],
        "private_rows_under_ignored_data": True,
        "public_receipt_metadata_only": True,
    }
    receipt = {**receipt_body, "receipt_id": canonical_sha256(receipt_body)}
    assert_metadata_only(receipt)
    _write_immutable_json(public_run_root(run_id) / "prepare-receipt.json", receipt)
    return receipt


def _validate_artifact_descriptor(descriptor: Mapping[str, Any]) -> Path:
    relative_path = descriptor.get("relative_path")
    if not isinstance(relative_path, str) or not relative_path:
        raise HumanReferenceError("artefact descriptor path is invalid")
    path = REPO_ROOT / relative_path
    if not path.is_file():
        raise HumanReferenceError(f"immutable artefact is missing: {path}")
    if (
        descriptor.get("sha256") != _sha256_file(path)
        or descriptor.get("bytes") != path.stat().st_size
    ):
        raise HumanReferenceError(f"immutable artefact digest or size drifted: {path}")
    return path


def load_reference_rows(
    manifest: Mapping[str, Any], *, split: str | None = None
) -> list[dict[str, Any]]:
    descriptor = manifest["private_artifacts"]["reference_rows"]
    path = _validate_artifact_descriptor(descriptor)
    value = _read_json(path)
    rows = value.get("rows")
    if not isinstance(rows, list):
        raise HumanReferenceError("private reference rows are invalid")
    if split is None:
        return [dict(row) for row in rows]
    if split not in {"development", "locked_test_candidate", "excluded_language"}:
        raise ValueError(f"invalid reference split: {split}")
    return [dict(row) for row in rows if row.get("split") == split]


def validate(manifest_file: Path) -> dict[str, Any]:
    manifest = _read_json(manifest_file)
    manifest_body = {key: value for key, value in manifest.items() if key != "manifest_id"}
    if manifest.get("manifest_id") != canonical_sha256(manifest_body):
        raise HumanReferenceError("manifest content address drifted")
    contract = manifest.get("contract")
    if not isinstance(contract, Mapping):
        raise HumanReferenceError("manifest contract is missing")
    run_id = canonical_sha256(contract)
    if manifest.get("run_id_value") != run_id:
        raise HumanReferenceError("manifest run ID differs from the frozen contract")
    if manifest_file.resolve() != manifest_path_for_run(run_id).resolve():
        raise HumanReferenceError("manifest path differs from its content-addressed run ID")
    if contract.get("schema_version") != SCHEMA_VERSION or contract.get("kind") != KIND:
        raise HumanReferenceError("manifest contract identity drifted")
    for descriptor in manifest.get("private_artifacts", {}).values():
        _validate_artifact_descriptor(descriptor)
    for descriptor in manifest.get("source_copies", {}).values():
        _validate_artifact_descriptor(descriptor)
    rows = load_reference_rows(manifest)
    descriptor = manifest["private_artifacts"]["reference_rows"]
    if descriptor.get("rows") != len(rows):
        raise HumanReferenceError("reference row conservation drifted")
    schema_copy = contract["sources"]["label_schema"]
    schema_path = REPO_ROOT / schema_copy["relative_path"]
    if _sha256_file(schema_path) != schema_copy["sha256"]:
        raise HumanReferenceError("label schema digest drifted")
    schema = load_label_schema(schema_path)
    for row in rows:
        exposure_reasons = row.get("exposure_reasons")
        if not isinstance(exposure_reasons, list) or exposure_reasons != sorted(
            set(exposure_reasons)
        ):
            raise HumanReferenceError("reference exposure reasons drifted")
        if not set(exposure_reasons).issubset(
            {
                "prior_machine_packet",
                "review50_overlap",
                "qualitative_first100_audit",
            }
        ):
            raise HumanReferenceError("reference exposure reason is unknown")
        expected_split = (
            "development"
            if row["language"] == "confident_english" and exposure_reasons
            else "locked_test_candidate"
            if row["language"] == "confident_english"
            else "excluded_language"
        )
        if row.get("split") != expected_split:
            raise HumanReferenceError("reference exposure-aware split drifted")
        expected_surface = {
            "target_text": row["target_text"],
            "submission_context": row["submission_context"],
            "parent_context": row["parent_context"],
        }
        if row.get("surface_sha256") != canonical_sha256(expected_surface):
            raise HumanReferenceError("reference surface hash drifted")
        expected_reference = {"language": row["language"], "label": row["label"]}
        if row.get("reference_sha256") != canonical_sha256(expected_reference):
            raise HumanReferenceError("reference label hash drifted")
        row_body = {key: value for key, value in row.items() if key != "row_sha256"}
        if row.get("row_sha256") != canonical_sha256(row_body):
            raise HumanReferenceError("reference row hash drifted")
        errors = list(Draft202012Validator(schema).iter_errors(row["label"]))
        if errors:
            raise HumanReferenceError("reference semantic label no longer validates")
        relevance = row["label"]["relevance"]
        target_stances = row["label"]["target_stances"]
        if (relevance == "material") != bool(target_stances):
            raise HumanReferenceError("reference cross-field semantic rule drifted")
    split_counts = Counter(row["split"] for row in rows)
    if manifest["summary"]["split_counts"] != {
        key: split_counts[key]
        for key in ("development", "locked_test_candidate", "excluded_language")
    }:
        raise HumanReferenceError("reference split conservation drifted")
    expected_split_counts = contract["split_policy"].get("expected_split_counts")
    if expected_split_counts is not None and manifest["summary"]["split_counts"] != dict(
        expected_split_counts
    ):
        raise HumanReferenceError("reference split counts differ from the frozen expectation")
    for artefact_name, summary_name in (
        ("quarantine_rows", "quarantine_rows"),
        ("review_rows", "review_valid_rows"),
    ):
        artefact_path = _validate_artifact_descriptor(
            manifest["private_artifacts"][artefact_name]
        )
        artefact_rows = _read_json(artefact_path).get("rows")
        if not isinstance(artefact_rows, list):
            raise HumanReferenceError(f"private {artefact_name} rows are invalid")
        if len(artefact_rows) != manifest["summary"][summary_name]:
            raise HumanReferenceError(f"private {artefact_name} conservation drifted")
    overlap_path = _validate_artifact_descriptor(
        manifest["private_artifacts"]["overlap_details"]
    )
    if manifest["overlap_summary"] != _overlap_summary(_read_json(overlap_path)):
        raise HumanReferenceError("overlap summary drifted")
    for source_name, copy_descriptor in manifest["source_copies"].items():
        if copy_descriptor["sha256"] != contract["sources"][source_name]["sha256"]:
            raise HumanReferenceError("source copy binding drifted")
    receipt_path = public_run_root(run_id) / "prepare-receipt.json"
    receipt = _read_json(receipt_path)
    receipt_body = {key: value for key, value in receipt.items() if key != "receipt_id"}
    if receipt.get("receipt_id") != canonical_sha256(receipt_body):
        raise HumanReferenceError("public receipt content address drifted")
    if receipt.get("manifest_id") != manifest["manifest_id"]:
        raise HumanReferenceError("public receipt manifest binding drifted")
    if receipt.get("manifest_sha256") != _sha256_file(manifest_file):
        raise HumanReferenceError("public receipt manifest hash drifted")
    assert_metadata_only(receipt)
    incomplete = list(PRIVATE_ROOT.rglob("*.incomplete")) + list(
        PUBLIC_ROOT.rglob("*.incomplete")
    )
    if incomplete:
        raise HumanReferenceError("incomplete human-reference outputs remain")
    return {
        "status": "valid",
        "run_id_value": run_id,
        "semantic_valid_rows": len(rows),
        "development_rows": split_counts["development"],
        "locked_test_candidate_rows": split_counts["locked_test_candidate"],
        "excluded_language_rows": split_counts["excluded_language"],
        "public_receipt_metadata_only": True,
    }


def _positive_int(value: str) -> int:
    parsed = int(value)
    if parsed <= 0 or not math.isfinite(parsed):
        raise argparse.ArgumentTypeError("must be a positive integer")
    return parsed


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("action", choices=("prepare", "validate"))
    parser.add_argument("--semantic-workbook", type=Path, default=DEFAULT_SEMANTIC_WORKBOOK)
    parser.add_argument("--review-workbook", type=Path, default=DEFAULT_REVIEW_WORKBOOK)
    parser.add_argument("--source-packet", type=Path, default=DEFAULT_SOURCE_PACKET)
    parser.add_argument("--schema-path", type=Path, default=DEFAULT_SCHEMA_PATH)
    parser.add_argument("--manifest-path", type=Path)
    parser.add_argument("--expected-semantic-rows", type=_positive_int, default=500)
    parser.add_argument("--expected-review-rows", type=_positive_int, default=50)
    parser.add_argument("--expected-source-packet-rows", type=_positive_int, default=150)
    args = parser.parse_args()
    if args.action == "prepare":
        result = prepare(
            semantic_workbook=args.semantic_workbook,
            review_workbook=args.review_workbook,
            source_packet=args.source_packet,
            schema_path=args.schema_path,
            expected_semantic_rows=args.expected_semantic_rows,
            expected_review_rows=args.expected_review_rows,
            expected_source_packet_rows=args.expected_source_packet_rows,
        )
    else:
        selected_manifest = args.manifest_path or manifest_path(
            semantic_workbook=args.semantic_workbook,
            review_workbook=args.review_workbook,
            source_packet=args.source_packet,
            schema_path=args.schema_path,
            expected_semantic_rows=args.expected_semantic_rows,
            expected_review_rows=args.expected_review_rows,
            expected_source_packet_rows=args.expected_source_packet_rows,
        )
        result = validate(selected_manifest)
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
