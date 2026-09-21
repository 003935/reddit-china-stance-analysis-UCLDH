from __future__ import annotations

import json
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from reddit_china_stance import human_reference_v1 as reference


def _semantic_values(
    sample_id: str,
    target_text: str,
    *,
    language: str = "confident_english",
    relevance: str | None = "material",
    target: str | None = "china_general",
    stance: str | None = "no_directed_stance",
    submission: str | None = None,
    parent: str | None = None,
) -> list[str | None]:
    return [
        sample_id,
        target_text,
        submission,
        parent,
        None,
        language,
        relevance,
        target,
        stance,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
    ]


def _write_semantic_workbook(path: Path, *, corrupt_header: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = reference.SEMANTIC_SHEET
    headers = list(reference.SEMANTIC_HEADERS)
    if corrupt_header:
        headers[1] = "Different target header"
    for column, value in enumerate(headers, start=1):
        sheet.cell(row=reference.SEMANTIC_HEADER_ROW, column=column, value=value)
    rows = {
        4: _semantic_values("H001", "old packet target"),
        5: _semantic_values(
            "H002",
            "review bridge target",
            submission="shared submission",
        ),
        6: _semantic_values("H003", "qualitatively inspected target"),
        104: _semantic_values("H004", "genuinely locked target"),
        105: _semantic_values(
            "H005", "non-English target", language="chinese"
        ),
        106: _semantic_values(
            "H006",
            "invalid target",
            relevance=None,
            target=None,
            stance=None,
        ),
        107: _semantic_values("H007", "still genuinely locked target"),
    }
    for row_index, values in rows.items():
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)
    workbook.save(path)


def _write_review_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = reference.REVIEW_SHEET
    sheet.append(list(reference.REVIEW_HEADERS))
    sheet.append(
        [
            "R01",
            "review bridge target",
            "shared submission",
            None,
            "relevant",
            "submission",
            "human reason",
            "none",
        ]
    )
    sheet.append(
        [
            "R02",
            "genuinely locked target",
            None,
            None,
            "relevant",
            None,
            "missing auxiliary context-use decision",
            "none",
        ]
    )
    workbook.save(path)


def _write_source_packet(path: Path) -> None:
    pq.write_table(
        pa.Table.from_pylist([{"record_id": "record-1", "target_text": "old packet target"}]),
        path,
    )


def _inputs(tmp_path: Path, *, corrupt_header: bool = False) -> dict[str, Path]:
    semantic = tmp_path / "semantic.xlsx"
    review = tmp_path / "review.xlsx"
    packet = tmp_path / "source-packet.parquet"
    schema = tmp_path / "human-reference-schema.json"
    _write_semantic_workbook(semantic, corrupt_header=corrupt_header)
    _write_review_workbook(review)
    _write_source_packet(packet)
    schema.write_bytes(reference.DEFAULT_SCHEMA_PATH.read_bytes())
    return {
        "semantic_workbook": semantic,
        "review_workbook": review,
        "source_packet": packet,
        "schema_path": schema,
    }


def _redirect_roots(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(reference, "REPO_ROOT", tmp_path)
    monkeypatch.setattr(reference, "PRIVATE_ROOT", tmp_path / "data/private-human-reference-v1")
    monkeypatch.setattr(reference, "PUBLIC_ROOT", tmp_path / "outputs/human-reference-v1")


def _prepare(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> tuple[dict, Path]:
    paths = _inputs(tmp_path)
    _redirect_roots(monkeypatch, tmp_path)
    receipt = reference.prepare(
        **paths,
        expected_semantic_rows=7,
        expected_review_rows=2,
        expected_source_packet_rows=1,
        expected_split_counts={
            "development": 4,
            "locked_test_candidate": 1,
            "excluded_language": 1,
        },
    )
    return receipt, reference.manifest_path_for_run(receipt["run_id_value"])


def test_prepare_freezes_exposure_aware_splits_and_private_rows(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    receipt, manifest_path = _prepare(monkeypatch, tmp_path)
    manifest = json.loads(manifest_path.read_text())
    rows = reference.load_reference_rows(manifest)
    by_id = {row["source_sample_id"]: row for row in rows}

    assert receipt["counts"]["split_counts"] == {
        "development": 4,
        "locked_test_candidate": 1,
        "excluded_language": 1,
    }
    assert by_id["H001"]["exposure_reasons"] == [
        "prior_machine_packet",
        "qualitative_first100_audit",
    ]
    assert by_id["H002"]["exposure_reasons"] == [
        "qualitative_first100_audit",
        "review50_overlap",
    ]
    assert by_id["H003"]["exposure_reasons"] == ["qualitative_first100_audit"]
    assert by_id["H004"]["split"] == "development"
    assert by_id["H004"]["exposure_reasons"] == ["review50_overlap"]
    assert by_id["H007"]["split"] == "locked_test_candidate"
    assert by_id["H007"]["exposure_reasons"] == []
    assert by_id["H005"]["split"] == "excluded_language"

    quarantine_path = tmp_path / manifest["private_artifacts"]["quarantine_rows"][
        "relative_path"
    ]
    quarantine = json.loads(quarantine_path.read_text())["rows"]
    assert len(quarantine) == 2
    assert {tuple(row["reason_codes"]) for row in quarantine} == {
        ("missing_relevance",),
        ("missing_or_invalid_context_used",),
    }
    assert reference.validate(manifest_path)["locked_test_candidate_rows"] == 1


def test_public_receipt_is_metadata_only_and_does_not_publish_reference_labels(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    receipt, _ = _prepare(monkeypatch, tmp_path)
    payload = json.dumps(receipt, sort_keys=True)
    assert '"label"' not in payload
    assert "target_text" not in payload
    assert "human reason" not in payload
    assert receipt["public_receipt_metadata_only"] is True


def test_prepare_rejects_header_drift(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    paths = _inputs(tmp_path, corrupt_header=True)
    _redirect_roots(monkeypatch, tmp_path)
    with pytest.raises(reference.HumanReferenceError, match="headers differ"):
        reference.prepare(
            **paths,
            expected_semantic_rows=7,
            expected_review_rows=2,
            expected_source_packet_rows=1,
            expected_split_counts=None,
        )


def test_split_gate_fails_before_any_private_or_public_output(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    paths = _inputs(tmp_path)
    _redirect_roots(monkeypatch, tmp_path)
    with pytest.raises(reference.HumanReferenceError, match="split counts differ"):
        reference.prepare(
            **paths,
            expected_semantic_rows=7,
            expected_review_rows=2,
            expected_source_packet_rows=1,
            expected_split_counts={
                "development": 3,
                "locked_test_candidate": 2,
                "excluded_language": 1,
            },
        )
    assert not reference.PRIVATE_ROOT.exists()
    assert not reference.PUBLIC_ROOT.exists()


def test_validate_rejects_private_row_tampering(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    _, manifest_path = _prepare(monkeypatch, tmp_path)
    manifest = json.loads(manifest_path.read_text())
    private_path = tmp_path / manifest["private_artifacts"]["reference_rows"]["relative_path"]
    private_path.write_text(private_path.read_text() + " ")
    with pytest.raises(reference.HumanReferenceError, match="digest or size drifted"):
        reference.validate(manifest_path)
